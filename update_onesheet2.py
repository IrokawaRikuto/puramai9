# -*- coding: utf-8 -*-
"""ワンシートのプレイ画面イメージを最新mockupに差し替え＋キャプション更新。"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE_TYPE

FN = "HEWワンシート企画書_プラマイナイン.pptx"
prs = Presentation(FN)
slide = prs.slides[0]

def find(sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    return None

# --- 既存の埋め込み画像(プレイ画面mockup)を削除 ---
for sh in list(slide.shapes):
    if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
        sh._element.getparent().remove(sh._element)
        print("removed picture", sh.shape_id)

# --- 本文ボックス[11]:概要＋キャプションを最新化 ---
box = find(11)
tf = box.text_frame
tf.word_wrap = True
tf.vertical_anchor = MSO_ANCHOR.TOP
tf.clear()

def set_para(p, text, size, bold, color, space_after=6):
    p.text = text
    p.alignment = PP_ALIGN.LEFT
    try: p.space_after = Pt(space_after)
    except Exception: pass
    for r in p.runs:
        r.font.size = Pt(size); r.font.bold = bold
        r.font.name = "Yu Gothic"; r.font.color.rgb = RGBColor(*color)

DARK = (0x22, 0x2A, 0x38)
ACC  = (0xC0, 0x4A, 0x2A)

set_para(tf.paragraphs[0],
         "数字を足し引きして、相手のHPを 0 か 200 の“端”へ押し出す 2〜4人対戦カードバトル。",
         13, True, DARK, space_after=5)
p1 = tf.add_paragraph()
set_para(p1,
         "コイントスの偶奇で、放った一手が「相手に命中」か「自分へ自爆」かが裏返る ── 運・計算・読み合いが一手に同居する心理戦。",
         11, False, DARK, space_after=8)
p2 = tf.add_paragraph()
set_para(p2,
         "▼ プレイ画面イメージ（共通デッキ115枚 ／ HP0〜200・非公開 ／ 前半2枚以上・後半1枚以上 ／ チャージで貯めるコスト制アイテム3種 ／ 1v1〜4人・レート対戦）",
         10, True, ACC, space_after=2)

# --- 最新mockupを再埋め込み ---
pic = slide.shapes.add_picture("mockup.png", Inches((7.5 - 6.90) / 2), Inches(5.50),
                               Inches(6.90), Inches(4.60))
pic.line.color.rgb = RGBColor(0x2C, 0x3A, 0x56)
pic.line.width = Pt(1.25)

prs.save(FN)
print("saved", FN)
