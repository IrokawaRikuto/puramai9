# -*- coding: utf-8 -*-
"""ワンシート企画書を最新化:本文を概要のみに差し替え＋プレイ画面イメージを埋め込み。"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

FN = "HEWワンシート企画書_プラマイナイン.pptx"
prs = Presentation(FN)
slide = prs.slides[0]

def find(sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    return None

# --- 1) 旧プレースホルダー矩形(画面イメージのダミー)を削除 ---
for sid in (1028, 1029, 1030, 1031):
    sh = find(sid)
    if sh is not None:
        sh._element.getparent().remove(sh._element)
        print("removed", sid)

# --- 2) メイン本文ボックス[11]:位置を下げて「概要のみ」に差し替え ---
box = find(11)
box.left   = Inches(0.13)
box.top    = Inches(4.50)
box.width  = Inches(7.24)
box.height = Inches(5.68)
tf = box.text_frame
tf.word_wrap = True
tf.vertical_anchor = MSO_ANCHOR.TOP
try:
    tf.margin_top = Inches(0.10); tf.margin_bottom = Inches(0.06)
    tf.margin_left = Inches(0.14); tf.margin_right = Inches(0.14)
except Exception:
    pass

# 既存段落クリア
tf.clear()

def set_para(p, text, size, bold, color, align=PP_ALIGN.LEFT, space_after=6):
    p.text = text
    p.alignment = align
    try: p.space_after = Pt(space_after)
    except Exception: pass
    for r in p.runs:
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = "Yu Gothic"
        r.font.color.rgb = RGBColor(*color)

DARK = (0x22, 0x2A, 0x38)
GREY = (0x5A, 0x66, 0x78)
ACC  = (0xC0, 0x4A, 0x2A)

p0 = tf.paragraphs[0]
set_para(p0, "数字を足し引きして、相手のHPを 0 か 200 の“端”へ押し出す 2人対戦カードバトル。",
         13, True, DARK, space_after=5)
p1 = tf.add_paragraph()
set_para(p1, "コイントスの偶奇で、放った一手が「相手に命中」か「自分へ自爆」かが裏返る ── 運・計算・読み合いが一手に同居する心理戦。",
         11, False, DARK, space_after=8)
p2 = tf.add_paragraph()
set_para(p2, "▼ プレイ画面イメージ（共通デッキ115枚 ／ HP0〜200・非公開 ／ 前半2枚以上・後半1枚以上 ／ ラウンド頭に当たり偶奇を公開するコイントス）",
         10, True, ACC, space_after=2)

# --- 3) プレイ画面イメージを埋め込み ---
img_w = Inches(6.90)
img_h = Inches(4.60)   # 1680x1120 ≒ 1.5:1
img_l = Inches((7.5 - 6.90) / 2)   # 中央寄せ
img_t = Inches(5.50)
pic = slide.shapes.add_picture("mockup.png", img_l, img_t, img_w, img_h)
# 枠線
ln = pic.line
ln.color.rgb = RGBColor(0x2C, 0x3A, 0x56)
ln.width = Pt(1.25)

prs.save(FN)
print("saved", FN)
